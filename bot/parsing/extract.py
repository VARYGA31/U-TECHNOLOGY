"""Чтение xlsx/xls/pdf/docx/csv/картинок + извлечение фото товаров.

Результат — текст с метками [ФОТО #n], словарь фото {n: bytes} и, для xlsx,
сам workbook со списком якорей картинок. Метки нужны ИИ, чтобы привязать
фото к позициям.
"""
from __future__ import annotations

import csv
import io
import os
import zipfile
from dataclasses import dataclass, field
from typing import Any

PHOTO_MARK = "[ФОТО #{}]"
MIN_PDF_IMAGE_PX = 120
MAX_SCAN_PAGES = 6
# Страница без текстового слоя, но с крупной картинкой — это лист каталога.
# Такие листы надо показать ИИ картинкой, иначе от них остаётся голая метка.
MAX_AI_IMAGES = 16
AI_IMAGE_PX = 1200
PAGE_TEXT_IS_EMPTY = 40


@dataclass
class Anchor:
    """Якорь картинки в xlsx."""
    sheet: str
    row: int        # с нуля
    col: int        # с нуля
    area: float     # площадь в пикселях — крупнейшая картинка считается фото товара
    n: int          # номер метки [ФОТО #n]


@dataclass
class Extracted:
    kind: str = ""
    text: str = ""
    photos: dict[int, bytes] = field(default_factory=dict)
    workbook: Any = None
    anchors: list[Anchor] = field(default_factory=list)
    # (номер метки [ФОТО #n] или 0, байты) — картинки, которые уходят в ИИ
    images_for_ai: list[tuple[int, bytes]] = field(default_factory=list)
    error: str = ""


def extract(path: str, filename: str | None = None) -> Extracted:
    name = (filename or os.path.basename(path)).lower()
    ext = os.path.splitext(name)[1]
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(path)
    if ext == ".xls":
        return _from_xls(path)
    if ext == ".pdf":
        return _from_pdf(path)
    if ext == ".docx":
        return _from_docx(path)
    if ext in (".csv", ".txt"):
        return _from_csv(path)
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"):
        return _from_image(path)
    return Extracted(kind="unknown", error=f"Формат {ext} не поддерживается")


# --------------------------------------------------------------- xlsx

def _from_xlsx(path: str) -> Extracted:
    import openpyxl

    # ВАЖНО: не read_only — иначе openpyxl не отдаст картинки
    wb = openpyxl.load_workbook(path, data_only=True)
    out = Extracted(kind="xlsx", workbook=wb)

    n = 0
    # картинки: байты читаем сразу — _data() отдаёт данные только один раз
    by_cell: dict[tuple[str, int, int], list[int]] = {}
    for ws in wb.worksheets:
        for img in getattr(ws, "_images", []):
            try:
                data = img._data()
            except Exception:
                continue
            if not data:
                continue
            n += 1
            frm = img.anchor._from
            area = float(getattr(img, "width", 0) or 0) * float(getattr(img, "height", 0) or 0)
            out.photos[n] = data
            out.anchors.append(Anchor(ws.title, frm.row, frm.col, area, n))
            by_cell.setdefault((ws.title, frm.row, frm.col), []).append(n)

    lines: list[str] = []
    for ws in wb.worksheets:
        if len(wb.worksheets) > 1:
            lines.append(f"### Лист: {ws.title}")
        merged = {}
        for rng in ws.merged_cells.ranges:
            if rng.max_row > rng.min_row:
                merged[(rng.min_row, rng.min_col)] = rng.max_row
        for row in ws.iter_rows():
            parts: list[str] = []
            for cell in row:
                marks = by_cell.get((ws.title, cell.row - 1, cell.column - 1), [])
                val = cell.value
                if val is None and not marks:
                    continue
                text = "" if val is None else str(val).strip()
                if not text and not marks:
                    continue
                chunk = f"{_col_letter(cell.column - 1)}: {text}"
                upto = merged.get((cell.row, cell.column))
                if upto:
                    chunk += f" (объединено до R{upto})"
                for m in marks:
                    chunk += " " + PHOTO_MARK.format(m)
                parts.append(chunk)
            if parts:
                lines.append(f"R{row[0].row} | " + " | ".join(parts))
    out.text = "\n".join(lines)
    return out


def _from_xls(path: str) -> Extracted:
    import xlrd

    book = xlrd.open_workbook(path)
    lines: list[str] = []
    for sheet in book.sheets():
        if book.nsheets > 1:
            lines.append(f"### Лист: {sheet.name}")
        for r in range(sheet.nrows):
            parts = []
            for c in range(sheet.ncols):
                v = sheet.cell_value(r, c)
                if v in ("", None):
                    continue
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                parts.append(f"{_col_letter(c)}: {v}")
            if parts:
                lines.append(f"R{r + 1} | " + " | ".join(parts))
    return Extracted(kind="xls", text="\n".join(lines))


def _col_letter(idx: int) -> str:
    s = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


# --------------------------------------------------------------- pdf

def _from_pdf(path: str) -> Extracted:
    import pdfplumber

    out = Extracted(kind="pdf")
    images: list[tuple[int, float, bytes]] = []  # (страница, y сверху, байты)
    try:
        images = _pdf_images(path)
    except Exception:
        images = []

    n = 0
    # список, а не словарь: на одной странице бывает несколько картинок с
    # одинаковой вертикальной позицией, в словаре они затирали бы друг друга
    marks: list[tuple[int, float, int]] = []
    images_with_marks: list[tuple[int, float, bytes, int]] = []
    for page_no, top, data in images:
        n += 1
        out.photos[n] = data
        marks.append((page_no, top, n))
        images_with_marks.append((page_no, top, data, n))

    lines: list[str] = []
    text_chars = 0
    page_chars: dict[int, int] = {}
    with pdfplumber.open(path) as pdf:
        pages = pdf.pages
        for pi, page in enumerate(pages):
            lines.append(f"--- Страница {pi + 1} ---")
            page_marks = sorted([(top, k) for p, top, k in marks if p == pi])
            rows = page.extract_text_lines() or []
            used = set()
            here = 0
            for line in rows:
                for top, k in page_marks:
                    if k not in used and top <= line["top"]:
                        lines.append(PHOTO_MARK.format(k))
                        used.add(k)
                text = (line.get("text") or "").strip()
                if text:
                    here += len(text)
                    lines.append(text)
            for top, k in page_marks:
                if k not in used:
                    lines.append(PHOTO_MARK.format(k))
            page_chars[pi] = here
            text_chars += here
    out.text = "\n".join(lines)

    if text_chars < 120:
        # текстового слоя нет вовсе — это скан, отдаём страницы ИИ картинками
        try:
            out.images_for_ai = [(0, data) for data in _pdf_render(path, MAX_SCAN_PAGES)]
            out.kind = "pdf-scan"
        except Exception:
            pass
        return out

    # Смешанный файл: таблица текстом + листы каталога картинками.
    # Страницы без текста отдаём ИИ изображениями, иначе от них в тексте
    # остаётся только метка [ФОТО #n], и привязать её не к чему.
    picture_pages = [(page_no, n, data) for page_no, _top, data, n in images_with_marks
                     if page_chars.get(page_no, 0) < PAGE_TEXT_IS_EMPTY]
    for _page_no, n, data in picture_pages[:MAX_AI_IMAGES]:
        small = _shrink(data, AI_IMAGE_PX)
        if small:
            out.images_for_ai.append((n, small))
    return out


def _shrink(data: bytes, max_px: int) -> bytes | None:
    from PIL import Image

    try:
        im = Image.open(io.BytesIO(data))
        im.load()
        im = im.convert("RGB")
        if max(im.size) > max_px:
            ratio = max_px / float(max(im.size))
            im = im.resize((max(1, int(im.width * ratio)),
                            max(1, int(im.height * ratio))), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=82, optimize=True)
        return buf.getvalue()
    except Exception:
        return None


def _pdf_images(path: str) -> list[tuple[int, float, bytes]]:
    import pypdfium2 as pdfium

    found: list[tuple[int, float, bytes]] = []
    pdf = pdfium.PdfDocument(path)
    try:
        for pi in range(len(pdf)):
            page = pdf[pi]
            height = page.get_height()
            try:
                objs = page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE])
            except Exception:
                continue
            for obj in objs:
                try:
                    pil = obj.get_bitmap(render=False).to_pil()
                except Exception:
                    continue
                if pil.width < MIN_PDF_IMAGE_PX or pil.height < MIN_PDF_IMAGE_PX:
                    continue
                try:
                    left, bottom, right, top = obj.get_pos()
                except Exception:
                    top = height
                buf = io.BytesIO()
                pil.convert("RGB").save(buf, "JPEG", quality=90)
                found.append((pi, float(height - top), buf.getvalue()))
    finally:
        pdf.close()
    return found


def _pdf_render(path: str, max_pages: int) -> list[bytes]:
    import pypdfium2 as pdfium

    out: list[bytes] = []
    pdf = pdfium.PdfDocument(path)
    try:
        for pi in range(min(len(pdf), max_pages)):
            pil = pdf[pi].render(scale=2).to_pil()
            buf = io.BytesIO()
            pil.convert("RGB").save(buf, "JPEG", quality=85)
            out.append(buf.getvalue())
    finally:
        pdf.close()
    return out


# --------------------------------------------------------------- docx

def _from_docx(path: str) -> Extracted:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    out = Extracted(kind="docx")
    doc = docx.Document(path)

    rels: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as zf:
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype:
                try:
                    rels[rel_id] = zf.read("word/" + rel.target_ref.lstrip("/").replace("word/", ""))
                except KeyError:
                    try:
                        rels[rel_id] = rel.target_part.blob
                    except Exception:
                        pass

    n = 0
    lines: list[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            blips = child.findall(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            for blip in blips:
                rid = blip.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                data = rels.get(rid)
                if data:
                    n += 1
                    out.photos[n] = data
                    text = (text + " " + PHOTO_MARK.format(n)).strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    out.text = "\n".join(lines)
    return out


# --------------------------------------------------------------- прочее

def _from_csv(path: str) -> Extracted:
    for enc in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            with open(path, newline="", encoding=enc) as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                except csv.Error:
                    dialect = csv.excel
                rows = [" | ".join(c.strip() for c in row if c is not None)
                        for row in csv.reader(fh, dialect)]
            return Extracted(kind="csv", text="\n".join(r for r in rows if r.strip(" |")))
        except (UnicodeDecodeError, OSError):
            continue
    return Extracted(kind="csv", error="Не удалось прочитать файл")


def _from_image(path: str) -> Extracted:
    with open(path, "rb") as fh:
        data = fh.read()
    return Extracted(kind="image", images_for_ai=[(0, data)])
