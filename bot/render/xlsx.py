"""Фирменный Excel — один в один с образцом kp_ru.xlsx."""
from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as PILImage

from ..config import Brand, settings
from ..models import CURRENCY_FORMAT, Quote, t
from . import labels as L
from .common import (col_width_to_px, fit_box, photo_bytes, prepare_photo,
                     pt_to_px, split_height, starts_with_option_word,
                     text_height_pt)

COL_WIDTH = {"no": 7.0, "name": 24.0, "photo": 38.0, "spec": 55.5,
             "qty": 10.0, "price": 18.0, "total": 18.0}
EXTRA_WIDTH = 22.0
MERGED_WIDTH = 50.0

ROW_LOGO = 79.0
ROW_TITLE = 38.0
ROW_SMALL = 19.0
ROW_HEADER = 69.0
ROW_BAR = 22.0
OPTION_MIN_PT = 50.0

THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _rgb(color: str) -> str:
    return "FF" + color.lstrip("#").upper()


def render_xlsx(quote: Quote, lang: str, out_path: str | Path,
                photo_dir: str | Path | None = None,
                brand: Brand | None = None) -> Path:
    brand = brand or settings.brand
    out_path = Path(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "1"

    columns = quote.active_columns()
    first_col = 2                       # таблица начинается с колонки B
    last_col = first_col + len(columns) - 1
    header_fill = PatternFill("solid", start_color=_rgb(brand.header_color))
    bar_fill = PatternFill("solid", start_color=_rgb(brand.bar_color))
    money_fmt = CURRENCY_FORMAT.get(quote.currency, '#,##0.00')

    # --- ширина колонок ---
    ws.column_dimensions["A"].width = 2.0
    for i, col in enumerate(columns):
        letter = get_column_letter(first_col + i)
        if col.key in COL_WIDTH:
            width = COL_WIDTH[col.key]
        elif col.merged:
            width = MERGED_WIDTH
        else:
            width = EXTRA_WIDTH
        ws.column_dimensions[letter].width = width

    def width_of(key: str) -> float:
        for i, col in enumerate(columns):
            if col.key == key:
                return ws.column_dimensions[get_column_letter(first_col + i)].width or 20.0
        return 20.0

    def col_index(key: str) -> int | None:
        for i, col in enumerate(columns):
            if col.key == key:
                return first_col + i
        return None

    def merge(r1: int, c1: int, r2: int, c2: int) -> None:
        if r2 > r1 or c2 > c1:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    # --- шапка документа ---
    ws.row_dimensions[1].height = ROW_LOGO
    _put_logo(ws, brand, first_col)

    row = 2
    ws.row_dimensions[row].height = ROW_TITLE
    cell = ws.cell(row=row, column=first_col, value=L.lb("doc_title", lang))
    cell.font = Font(name=brand.font_title, size=20, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    merge(row, first_col, row, last_col)
    row += 1

    if quote.client:
        ws.row_dimensions[row].height = ROW_SMALL
        cell = ws.cell(row=row, column=first_col,
                       value=L.lb("for_client", lang).format(quote.client))
        cell.font = Font(name=brand.font_title, size=14, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        merge(row, first_col, row, last_col)
        row += 1

    ws.row_dimensions[row].height = ROW_SMALL
    cell = ws.cell(row=row, column=first_col,
                   value=f"{L.lb('date', lang)}: {L.date_str(quote.created, lang)}")
    cell.font = Font(name=brand.font_title, size=14, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    merge(row, first_col, row, last_col)
    row += 1

    # --- строка-шапка таблицы ---
    header_row = row
    ws.row_dimensions[header_row].height = ROW_HEADER
    for i, col in enumerate(columns):
        if col.key in ("price", "total"):
            title = L.money_header(col.key, lang, quote.currency, quote.incoterms)
        else:
            title = t(col.title, lang) or L.BASE_TITLES.get(col.key, {}).get(lang, "")
            if col.is_money:
                title = f"{title}\n({L.currency_label(quote.currency)})"
        cell = ws.cell(row=header_row, column=first_col + i, value=title)
        cell.font = Font(name=brand.font_title, size=12, bold=True)
        cell.fill = header_fill
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1

    body_start = row
    spec_width = width_of("spec")
    name_width = width_of("name")

    # --- позиции ---
    for index, product in enumerate(quote.products, 1):
        spec_text = t(product.spec, lang)
        name_text = t(product.name, lang)

        need = max(
            text_height_pt(spec_text, spec_width) if quote.has_column("spec") else 0.0,
            text_height_pt(name_text, name_width, line_pt=19.0),
            90.0 if product.photo and quote.has_column("photo") else 24.0,
        )
        main_heights = split_height(need)
        main_start = row
        main_end = row + len(main_heights) - 1
        for offset, height in enumerate(main_heights):
            ws.row_dimensions[main_start + offset].height = height

        opt_rows: list[int] = []
        opt_start = main_end + 1
        for option in product.options:
            opt_row = opt_start + len(opt_rows)
            ws.row_dimensions[opt_row].height = max(
                OPTION_MIN_PT, text_height_pt(t(option.text, lang), spec_width))
            opt_rows.append(opt_row)
        pos_end = opt_rows[-1] if opt_rows else main_end

        # все опции уже начинаются со слова «Опция» → отдельная подпись не нужна
        label_needed = bool(opt_rows) and not all(
            starts_with_option_word(t(o.text, lang)) for o in product.options)

        for i, col in enumerate(columns):
            c = first_col + i
            if col.merged:
                continue  # такие колонки заполняем один раз на всё КП

            if col.key == "no":
                cell = ws.cell(row=main_start, column=c, value=index)
                cell.font = Font(name=brand.font_title, size=12, bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
                merge(main_start, c, pos_end, c)

            elif col.key == "name":
                cell = ws.cell(row=main_start, column=c, value=name_text)
                cell.font = Font(name=brand.font_text, size=14)
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
                if label_needed:
                    merge(main_start, c, main_end, c)
                    label = ws.cell(row=opt_rows[0], column=c, value=L.lb("option", lang))
                    label.font = Font(name=brand.font_text, size=14)
                    label.alignment = Alignment(horizontal="center", vertical="center",
                                                wrap_text=True)
                    merge(opt_rows[0], c, opt_rows[-1], c)
                else:
                    merge(main_start, c, pos_end, c)

            elif col.key == "photo":
                merge(main_start, c, pos_end, c)
                ws.cell(row=main_start, column=c).alignment = Alignment(
                    horizontal="center", vertical="center")
                _put_photo(ws, quote, product, photo_dir, main_start, c,
                           width_of("photo"),
                           sum(main_heights) + sum(
                               ws.row_dimensions[r].height or 15 for r in opt_rows))

            elif col.key == "spec":
                cell = ws.cell(row=main_start, column=c, value=spec_text)
                cell.font = Font(name=brand.font_text, size=12)
                cell.alignment = Alignment(horizontal="left", vertical="center",
                                           wrap_text=True)
                merge(main_start, c, main_end, c)
                for opt_row, option in zip(opt_rows, product.options):
                    text = t(option.text, lang)
                    ocell = ws.cell(row=opt_row, column=c, value=text)
                    ocell.font = Font(name=brand.font_text, size=12)
                    ocell.alignment = Alignment(horizontal="left", vertical="center",
                                                wrap_text=True)

            elif col.key == "qty":
                _num(ws, main_start, c, product.qty, brand, None)
                merge(main_start, c, main_end, c)
                for opt_row, option in zip(opt_rows, product.options):
                    _num(ws, opt_row, c, option.qty, brand, None)

            elif col.key == "price":
                _num(ws, main_start, c, quote.price(product), brand, money_fmt)
                merge(main_start, c, main_end, c)
                for opt_row, option in zip(opt_rows, product.options):
                    _num(ws, opt_row, c, quote.option_price(product, option), brand, money_fmt)

            elif col.key == "total":
                _num(ws, main_start, c, quote.total(product), brand, money_fmt)
                merge(main_start, c, main_end, c)
                for opt_row, option in zip(opt_rows, product.options):
                    _num(ws, opt_row, c, quote.option_total(product, option), brand, money_fmt)

            elif col.is_money:
                value = quote.money_value(product, col)
                if value is not None:
                    _num(ws, main_start, c, value, brand, money_fmt)
                merge(main_start, c, pos_end, c)

            else:
                cell = ws.cell(row=main_start, column=c, value=t(product.texts.get(col.key), lang))
                cell.font = Font(name=brand.font_text, size=12)
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
                merge(main_start, c, pos_end, c)

        row = pos_end + 1

    body_end = row - 1

    # --- колонки с одним значением на всё КП ---
    for i, col in enumerate(columns):
        if not col.merged:
            continue
        c = first_col + i
        cell = ws.cell(row=body_start, column=c, value=t(quote.merged.get(col.key), lang))
        cell.font = Font(name=brand.font_text, size=14)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        merge(body_start, c, body_end, c)

    # --- рамки по всей таблице ---
    for r in range(header_row, body_end + 1):
        for c in range(first_col, last_col + 1):
            ws.cell(row=r, column=c).border = BORDER

    # --- итоги ---
    total_col = col_index("total") or last_col
    label_col = col_index("name") or first_col
    if quote.show_total:
        row = _totals(ws, quote, lang, brand, row, first_col, label_col,
                      total_col, last_col, money_fmt)

    # --- фирменная полоса ---
    ws.row_dimensions[row].height = ROW_BAR
    for c in range(first_col, last_col + 1):
        ws.cell(row=row, column=c).fill = bar_fill
    merge(row, first_col, row, last_col)
    bar_row = row
    row += 2

    # --- условия и сноска ---
    row = _terms(ws, quote, lang, brand, row, first_col, last_col)
    footer = brand.footer(lang)
    if footer:
        cell = ws.cell(row=row, column=first_col, value=footer)
        cell.font = Font(name=brand.font_text, size=10, italic=True)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        merge(row, first_col, row, last_col)
        ws.row_dimensions[row].height = 28.0

    # --- печать ---
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{header_row}"
    ws.print_area = (f"A1:{get_column_letter(last_col)}{bar_row}")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=first_col)
    ws.sheet_view.zoomScale = 70

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _num(ws, row: int, col: int, value: float, brand: Brand, fmt: str | None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=brand.font_title, size=14)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fmt:
        cell.number_format = fmt


def _totals(ws, quote: Quote, lang: str, brand: Brand, row: int, first_col: int,
            label_col: int, total_col: int, last_col: int, money_fmt: str) -> int:
    def line(label: str, amount: float, bold: bool) -> None:
        nonlocal row
        ws.row_dimensions[row].height = 24.0
        cell = ws.cell(row=row, column=first_col, value=label)
        cell.font = Font(name=brand.font_title, size=12, bold=bold)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        if total_col - 1 >= first_col:
            ws.merge_cells(start_row=row, start_column=first_col,
                           end_row=row, end_column=total_col - 1)
        value = ws.cell(row=row, column=total_col, value=amount)
        value.font = Font(name=brand.font_title, size=12, bold=bold)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.number_format = money_fmt
        for c in range(first_col, last_col + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    line(L.lb("subtotal", lang), quote.subtotal(), False)
    for ln in quote.lines:
        line(t(ln.label, lang), quote.line_amount(ln), False)
    if quote.lines:
        line(L.lb("grand_total", lang), quote.grand_total(), True)
    return row


def _terms(ws, quote: Quote, lang: str, brand: Brand, row: int,
           first_col: int, last_col: int) -> int:
    order = [("payment", "payment"), ("lead", "lead"),
             ("warranty", "warranty"), ("extra", "extra")]
    if quote.incoterms:
        cell = ws.cell(row=row, column=first_col,
                       value=f"{L.lb('incoterms', lang)} {quote.incoterms}")
        cell.font = Font(name=brand.font_text, size=12, bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=first_col,
                       end_row=row, end_column=last_col)
        ws.row_dimensions[row].height = 22.0
        row += 1
    for key, label_key in order:
        text = t(quote.terms.get(key), lang)
        if not text:
            continue
        cell = ws.cell(row=row, column=first_col,
                       value=f"{L.lb(label_key, lang)} {text}")
        cell.font = Font(name=brand.font_text, size=12)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=first_col,
                       end_row=row, end_column=last_col)
        ws.row_dimensions[row].height = max(
            22.0, text_height_pt(text, 110.0, line_pt=15.0))
        row += 1
    return row + 1


def _put_logo(ws, brand: Brand, col: int) -> None:
    path = brand.logo_path
    if not path.exists():
        return
    try:
        pil = PILImage.open(path)
        pil.load()
    except Exception:
        return
    box_h = pt_to_px(ROW_LOGO) - 12
    box_w = col_width_to_px(COL_WIDTH.get("no", 7.0) + 24.0)
    w, h = fit_box(pil.width, pil.height, box_w, box_h)
    img = XLImage(str(path))
    img.width, img.height = int(w), int(h)
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(4),
                           row=0, rowOff=pixels_to_EMU(6)),
        ext=XDRPositiveSize2D(pixels_to_EMU(int(w)), pixels_to_EMU(int(h))),
    )
    ws.add_image(img)


def _put_photo(ws, quote: Quote, product, photo_dir, row: int, col: int,
               col_chars: float, height_pt: float) -> None:
    data = photo_bytes(photo_dir, product.photo)
    if not data:
        return
    data = prepare_photo(data)
    try:
        pil = PILImage.open(io.BytesIO(data))
        pil.load()
    except Exception:
        return
    box_w = col_width_to_px(col_chars) - 10
    box_h = pt_to_px(height_pt) - 10
    if box_w <= 10 or box_h <= 10:
        return
    w, h = fit_box(pil.width, pil.height, box_w, box_h)
    img = XLImage(io.BytesIO(data))
    img.width, img.height = int(w), int(h)
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(int((box_w - w) / 2) + 5),
                           row=row - 1, rowOff=pixels_to_EMU(int((box_h - h) / 2) + 5)),
        ext=XDRPositiveSize2D(pixels_to_EMU(int(w)), pixels_to_EMU(int(h))),
    )
    ws.add_image(img)
